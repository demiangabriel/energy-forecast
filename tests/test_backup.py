"""Copie de rezervă a configurării parcurilor (secțiunea 16) — Excel și CSV."""
from __future__ import annotations

from openpyxl import load_workbook
import io

from app.core.backup import backup_parcuri_csv, backup_parcuri_xlsx
from tests.helpers import parc_eolian, parc_pv


def test_xlsx_contine_un_rand_per_parc_cu_antet(repo):
    repo.salveaza_parc(parc_pv(cod="PV-A"))
    repo.salveaza_parc(parc_eolian(cod="EO-B"))

    wb = load_workbook(io.BytesIO(backup_parcuri_xlsx(repo)))
    ws = wb["Parcuri"]
    coduri = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert coduri == {"PV-A", "EO-B"}
    assert ws.cell(row=1, column=1).value == "cod"


def test_csv_are_coloanele_specifice_tehnologiei(repo):
    repo.salveaza_parc(parc_pv(cod="PV-A"))
    continut = backup_parcuri_csv(repo)
    antet, rand_pv = continut.splitlines()[0], continut.splitlines()[1]
    assert "putere_dc_kwp" in antet
    coloane = antet.split(";")
    valori = rand_pv.split(";")
    assert valori[coloane.index("putere_dc_kwp")] != ""
