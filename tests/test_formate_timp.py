"""Armonizarea formatului de timestamp în toate fișierele produse de platformă
(`app/core/momente.py`).

Problema rezolvată: importul cere din v8 decalaj de fus orar explicit, dar
exportul propriu scria ora locală FĂRĂ decalaj. Un fișier produs de platformă
nu trecea, deci, prin importul platformei — iar orele din nopțile schimbării
orei rămâneau ambigue exact ca înainte.
"""
from __future__ import annotations

import csv
import io

import pytest

from app.core.csv_import import analizeaza_fisier
from app.core.momente import (FORMAT_TIMP_IMPLICIT, FORMATE_TIMP, format_timp_valid,
                               formateaza_moment)
from app.models import OraPrognoza, Rulare, new_id
from tests.helpers import parc_pv

TS = "2026-06-01T09:00:00+00:00"   # = 12:00 ora României, vara


def test_cele_trei_formate_scriu_acelasi_moment_diferit():
    assert formateaza_moment(TS, "Europe/Bucharest", "iso_local") == "2026-06-01T12:00:00+03:00"
    assert formateaza_moment(TS, "Europe/Bucharest", "iso_utc") == "2026-06-01T09:00:00+00:00"
    assert formateaza_moment(TS, "Europe/Bucharest", "local_simplu") == "2026-06-01 12:00"


def test_implicitul_e_cel_cu_decalaj():
    assert FORMAT_TIMP_IMPLICIT == "iso_local"
    assert formateaza_moment(TS, "Europe/Bucharest") == "2026-06-01T12:00:00+03:00"


def test_format_necunoscut_cade_pe_implicit_fara_eroare():
    """Un parametru greșit în URL nu trebuie să producă o eroare, ci formatul
    sigur — altfel o descărcare eșuează din cauza unei litere greșite."""
    assert format_timp_valid("aiurea") == FORMAT_TIMP_IMPLICIT
    assert format_timp_valid(None) == FORMAT_TIMP_IMPLICIT
    assert format_timp_valid("iso_utc") == "iso_utc"


@pytest.mark.parametrize("cheie", list(FORMATE_TIMP))
def test_fiecare_format_are_eticheta_exemplu_si_explicatie(cheie):
    """Lista e afișată direct în interfață — un format fără explicație ar pune
    operatorul să aleagă pe ghicite."""
    info = FORMATE_TIMP[cheie]
    assert info["eticheta"] and info["exemplu"] and info["explicatie"]


# --------------------------------------------------------------- export CSV


def _export(client, parc, **param) -> list[dict]:
    r = client.get(f"/parcuri/{parc.id}/export.csv", params=param)
    assert r.status_code == 200, r.status_code
    return list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))


def _parc_cu_rulare(repo):
    parc = parc_pv(cod="FMT-EXP", fus_orar="Europe/Bucharest")
    repo.salveaza_parc(parc)
    repo.salveaza_rulare(Rulare(
        id=new_id(), parc_id=parc.id, declansata_la=TS, tip="manuala", stare="reusita",
        ore=[OraPrognoza(timestamp_utc=TS, putere_ac_fizic_kw=400.0, putere_finala_kw=410.0,
                          fractie_incarcare=0.5)]))
    return parc


def test_exportul_csv_scrie_implicit_cu_decalaj(client, repo):
    parc = _parc_cu_rulare(repo)
    randuri = _export(client, parc)
    assert randuri[0]["timestamp"] == "2026-06-01T12:00:00+03:00"
    # coloana UTC brută rămâne mereu, ca sursă de adevăr
    assert randuri[0]["timestamp_utc"] == TS


def test_exportul_csv_respecta_formatul_ales(client, repo):
    parc = _parc_cu_rulare(repo)
    assert _export(client, parc, format_timp="iso_utc")[0]["timestamp"] == "2026-06-01T09:00:00+00:00"
    assert _export(client, parc, format_timp="local_simplu")[0]["timestamp"] == "2026-06-01 12:00"


def test_exportul_csv_ignora_un_format_gresit(client, repo):
    parc = _parc_cu_rulare(repo)
    assert _export(client, parc, format_timp="nuexista")[0]["timestamp"] == "2026-06-01T12:00:00+03:00"


def test_fisierul_exportat_trece_prin_propriul_import(client, repo):
    """Testul care justifică armonizarea: ce scrie platforma trebuie să poată
    intra înapoi în platformă, fără conversii manuale și fără bifa de
    compatibilitate pentru formatul vechi."""
    parc = _parc_cu_rulare(repo)
    randuri = _export(client, parc)

    csv_reimport = "timestamp;kWh\n" + "\n".join(
        f"{r['timestamp']};{r['putere_kw']}" for r in randuri)
    raport = analizeaza_fisier(csv_reimport.encode("utf-8"), parc)

    assert raport.valid, raport.erori
    assert raport.randuri_valide[0].timestamp_utc == TS, \
        "momentul trebuie să se întoarcă identic după export + import"


def test_formatul_vechi_exportat_NU_trece_prin_import(client, repo):
    """Perechea testului de mai sus — arată de ce vechiul format implicit era
    o problemă, nu o preferință de stil."""
    parc = _parc_cu_rulare(repo)
    randuri = _export(client, parc, format_timp="local_simplu")
    csv_reimport = "timestamp;kWh\n" + "\n".join(
        f"{r['timestamp']};{r['putere_kw']}" for r in randuri)

    raport = analizeaza_fisier(csv_reimport.encode("utf-8"), parc)
    assert not raport.valid
    assert any("decalaj" in e for e in raport.erori)


# --------------------------------------------------------------- raport lunar


def test_raportul_lunar_foloseste_acelasi_format(repo):
    from openpyxl import load_workbook

    from app.core.reports import genereaza_raport_lunar
    from app.models import ProductieOrara

    parc = parc_pv(cod="FMT-RAP", fus_orar="Europe/Bucharest")
    repo.salveaza_parc(parc)
    repo.salveaza_productie([ProductieOrara(parc_id=parc.id, timestamp_utc=TS, energie_kwh=400.0)])

    for fmt, asteptat in [("iso_local", "2026-06-01T12:00:00+03:00"),
                           ("iso_utc", "2026-06-01T09:00:00+00:00"),
                           ("local_simplu", "2026-06-01 12:00")]:
        wb = load_workbook(io.BytesIO(genereaza_raport_lunar(repo, parc, 2026, 6, format_timp=fmt)))
        assert wb["Date orare"].cell(row=2, column=1).value == asteptat, fmt


def test_raportul_lunar_pastreaza_coloana_utc_bruta(repo):
    from openpyxl import load_workbook

    from app.core.reports import genereaza_raport_lunar
    from app.models import ProductieOrara

    parc = parc_pv(cod="FMT-RAP2", fus_orar="Europe/Bucharest")
    repo.salveaza_parc(parc)
    repo.salveaza_productie([ProductieOrara(parc_id=parc.id, timestamp_utc=TS, energie_kwh=400.0)])

    wb = load_workbook(io.BytesIO(genereaza_raport_lunar(repo, parc, 2026, 6,
                                                          format_timp="local_simplu")))
    assert wb["Date orare"].cell(row=2, column=8).value == TS, \
        "sursa de adevăr UTC rămâne în raport indiferent de formatul ales"
