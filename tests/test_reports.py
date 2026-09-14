"""Raportul lunar Excel — inclusiv regresia pentru formulele care trimiteau
spre celula greșită și lăsau indicatorii principali GOI în raportul livrat."""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.models import OraPrognoza, ProductieOrara, Rulare, new_id
from tests.helpers import parc_pv


def _pregateste(repo, ore: int = 48, prognozat: float = 100.0, realizat: float = 90.0):
    parc = parc_pv()
    repo.salveaza_parc(parc)
    momente = [f"2026-06-{(h // 24) + 1:02d}T{h % 24:02d}:00:00+00:00" for h in range(ore)]
    repo.salveaza_productie([
        ProductieOrara(parc_id=parc.id, timestamp_utc=ts, energie_kwh=realizat) for ts in momente])
    repo.salveaza_rulare(Rulare(
        id=new_id(), parc_id=parc.id, declansata_la="2026-05-31T06:00:00+00:00",
        tip="programata", stare="reusita",
        ore=[OraPrognoza(timestamp_utc=ts, putere_finala_kw=prognozat, ore_in_avans=24.0)
             for ts in momente]))
    return parc


def _workbook(continut: bytes):
    return load_workbook(io.BytesIO(continut))


def test_raportul_lunar_se_genereaza(repo):
    from app.core.reports import genereaza_raport_lunar
    parc = _pregateste(repo)
    wb = _workbook(genereaza_raport_lunar(repo, parc, 2026, 6))
    assert {"Date orare", "Sumar", "Verificare"}.issubset(set(wb.sheetnames))


def test_REGRESIE_formulele_trimit_spre_celula_corecta(repo):
    """BUG ISTORIC: nMAE și bias împărțeau la C3, dar puterea instalată era
    scrisă în B3 (coloana A = etichetă, coloana B = valoare). Cum formulele
    erau împachetate în IFERROR, eroarea de referință nu se vedea: cei doi
    indicatori principali apăreau pur și simplu GOI în raportul trimis
    clientului, lună de lună.

    Testul verifică structural că fiecare formulă indică o celulă care chiar
    conține baza de normalizare.
    """
    from app.core.reports import genereaza_raport_lunar
    parc = _pregateste(repo)
    wb = _workbook(genereaza_raport_lunar(repo, parc, 2026, 6))
    ws = wb["Sumar"]

    etichete = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    rand_baza = etichete["Bază de normalizare a indicatorilor (MW)"]
    assert isinstance(ws.cell(row=rand_baza, column=2).value, (int, float)), \
        "baza de normalizare trebuie să fie o valoare numerică, în coloana B"

    referinta_asteptata = f"$B${rand_baza}"
    for eticheta in ["Eroare medie absolută — nMAE (% din puterea de referință)",
                      "Eroare pătratică medie — nRMSE (% din puterea de referință)",
                      "Abatere sistematică — bias (%, cu semn)"]:
        formula = ws.cell(row=etichete[eticheta], column=2).value
        assert isinstance(formula, str) and formula.startswith("="), f"{eticheta} nu e formulă"
        assert referinta_asteptata in formula, (
            f"'{eticheta}' se raportează la {referinta_asteptata}? Formula găsită: {formula}")


def test_etichetele_de_acuratete_sunt_neambigue(repo):
    """Două rânduri cu exact aceeași etichetă („nMAE (%)") în aceeași foaie
    lăsau cititorul să ghicească la ce se raportează fiecare. Baza de
    normalizare trebuie să fie vizibilă în etichetă."""
    from app.core.reports import genereaza_raport_lunar
    parc = _pregateste(repo)
    wb = _workbook(genereaza_raport_lunar(repo, parc, 2026, 6))
    ws = wb["Sumar"]
    etichete = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    etichete_reale = [e for e in etichete if e]
    assert len(etichete_reale) == len(set(etichete_reale)), \
        f"etichete duplicate în foaia Sumar: {[e for e in etichete_reale if etichete_reale.count(e) > 1]}"


def test_formula_de_energie_indica_randurile_corecte(repo):
    from app.core.reports import genereaza_raport_lunar
    parc = _pregateste(repo)
    wb = _workbook(genereaza_raport_lunar(repo, parc, 2026, 6))
    ws = wb["Sumar"]
    etichete = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    rand_prog = etichete["Energie totală prognozată (MWh)"]
    rand_real = etichete["Energie totală realizată (MWh)"]
    formula_abatere = ws.cell(row=etichete["Abatere de energie (%)"], column=2).value
    assert f"B{rand_prog}" in formula_abatere and f"B{rand_real}" in formula_abatere


def test_coloana_de_eroare_cu_semn_exista_si_e_formula(repo):
    """Biasul se calculează acum dintr-o coloană explicită de eroare cu semn,
    definită doar pe rândurile cu AMBELE valori — înainte suma prognozei
    includea și orele fără realizat, deci se împărțeau două mulțimi diferite."""
    from app.core.reports import genereaza_raport_lunar
    parc = _pregateste(repo)
    wb = _workbook(genereaza_raport_lunar(repo, parc, 2026, 6))
    ws = wb["Date orare"]
    assert ws.cell(row=1, column=5).value == "Eroare cu semn (MW)"
    formula = ws.cell(row=2, column=5).value
    assert formula.startswith("=IF(AND(B2<>") and "B2-C2" in formula


def test_foaia_de_verificare_contine_skill_si_excluderi(repo):
    from app.core.reports import genereaza_raport_lunar
    parc = _pregateste(repo)
    wb = _workbook(genereaza_raport_lunar(repo, parc, 2026, 6))
    ws = wb["Verificare"]
    texte = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert any("SKILL" in (t or "").upper() for t in texte)
    assert any("noapte" in (t or "").lower() for t in texte)
    assert any("calitate" in (t or "").lower() for t in texte)


def test_raportul_pentru_o_luna_fara_date_nu_crapa(repo):
    from app.core.reports import genereaza_raport_lunar
    parc = parc_pv()
    repo.salveaza_parc(parc)
    continut = genereaza_raport_lunar(repo, parc, 2020, 1)
    assert continut[:2] == b"PK"


def test_raportul_de_verificare_pentru_interfata(repo):
    from app.core.reports import genereaza_raport_verificare
    parc = _pregateste(repo, ore=72)
    rezultat = genereaza_raport_verificare(repo, parc, zile=3650)
    assert rezultat["disponibil"] is True
    assert rezultat["raport"].indicatori.numar_ore > 0
    assert rezultat["baza_normalizare_kw"] == parc.config_pv.putere_ac_kw


def test_coloana_de_timp_gestioneaza_corect_ora_de_iarna(repo):
    """R-1502: raportul lunar are o coloană de timp în fusul ales la
    descărcare, pe lângă timestamp-ul UTC brut. Octombrie 2026 conține
    tranziția la ora de iarnă (25 octombrie).

    Din v8, formatul implicit include decalajul explicit — ceea ce face ca ora
    locală 03:00, care se petrece de DOUĂ ori în acea noapte, să apară ca două
    valori DISTINCTE (+03:00 și +02:00). Înainte, fără decalaj, cele două erau
    scrise identic, iar cititorul raportului nu avea cum să le deosebească."""
    from collections import Counter

    from app.core.reports import genereaza_raport_lunar
    parc = parc_pv(fus_orar="Europe/Bucharest")
    repo.salveaza_parc(parc)
    momente = [f"2026-10-{d:02d}T{h:02d}:00:00+00:00" for d in (24, 25, 26) for h in range(24)]
    repo.salveaza_productie([
        ProductieOrara(parc_id=parc.id, timestamp_utc=ts, energie_kwh=50.0) for ts in momente])

    wb = _workbook(genereaza_raport_lunar(repo, parc, 2026, 10, fus_orar="Europe/Bucharest"))
    ws = wb["Date orare"]
    assert ws.cell(row=1, column=1).value == "Timestamp (Europe/Bucharest)"
    assert ws.cell(row=1, column=8).value == "Timestamp (UTC)"

    valori = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert all(valori), "toate rândurile trebuie să aibă coloana de timp completată"
    assert max(Counter(valori).values()) == 1, \
        "cu decalaj explicit, cele două ore locale 03:00 din noaptea tranziției sunt distincte"
    # exact perechea care înainte se scria identic
    assert "2026-10-25T03:00:00+03:00" in valori
    assert "2026-10-25T03:00:00+02:00" in valori

    # formatul vechi, cerut explicit, reproduce ambiguitatea — dovada că
    # avertismentul din interfață nu e teoretic
    ws_vechi = _workbook(genereaza_raport_lunar(repo, parc, 2026, 10, fus_orar="Europe/Bucharest",
                                                 format_timp="local_simplu"))["Date orare"]
    valori_vechi = [ws_vechi.cell(row=r, column=1).value for r in range(2, ws_vechi.max_row + 1)]
    assert max(Counter(valori_vechi).values()) == 2

    # alegerea unui alt fus la descărcare schimbă doar antetul/coloana A
    wb_utc = _workbook(genereaza_raport_lunar(repo, parc, 2026, 10, fus_orar="UTC"))
    ws_utc = wb_utc["Date orare"]
    assert ws_utc.cell(row=1, column=1).value == "Timestamp (UTC)"
    assert ws_utc.cell(row=2, column=1).value == ws_utc.cell(row=2, column=8).value


def test_raportul_de_verificare_fara_date_este_marcat_indisponibil(repo):
    from app.core.reports import genereaza_raport_verificare
    parc = parc_pv()
    repo.salveaza_parc(parc)
    assert genereaza_raport_verificare(repo, parc)["disponibil"] is False
