"""
Rapoarte, comparații și trasabilitate — secțiunea 12 (R-1201-1204).

Raportul lunar (R-1202) e un fișier de calcul REAL: indicatorii de sumar
sunt FORMULE Excel care citesc din foaia de date, nu valori fixe — dacă
operatorul editează datele brute, indicatorii se recalculează singuri, așa
cum cere explicit specificația.

--------------------------------------------------------------------------
CORECȚII ȘI COMPLETĂRI v2 (vezi docs/DEBUGGING.md):

 * BUG CORECTAT: formulele pentru nMAE și bias împărțeau la celula C3, în
   timp ce puterea instalată era scrisă în B3. Rezultatul: ambii indicatori
   principali ai raportului lunar apăreau GOI (IFERROR ascundea eroarea de
   referință). Raportul livrat clientului nu conținea, practic, niciun
   indicator de acuratețe.

 * BUG CORECTAT: biasul împărțea suma diferențelor la numărul de ore cu
   producție realizată, dar suma prognozei includea și orele fără realizat —
   două mulțimi diferite de rânduri. Acum există o coloană explicită de
   eroare cu semn, calculată doar pe rândurile cu ambele valori.

 * COMPLETARE: raportul conține acum și o foaie de verificare cu scorul de
   skill față de prognoza de referință, indicatori pe orizonturi și
   normalizare dublă (la puterea instalată ȘI la producția medie) — conform
   consensului din literatură. Un "nMAE" fără bază de normalizare declarată
   nu e apărabil într-o discuție comercială.
"""
from __future__ import annotations

import calendar
import io
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.forecast_engine import putere_instalata_kw, putere_normalizare_kw
from app.core.formatare import kw_in_mw
from app.core.momente import iso_utc
from app.models import Parc

FORMAT_MW = "0.000"

ANTET_FILL = PatternFill(start_color="1F6F50", end_color="1F6F50", fill_type="solid")
ANTET_FONT = Font(color="FFFFFF", bold=True)
SECTIUNE_FONT = Font(bold=True, size=12)


def _prognoza_pe_ora(repo, parc_id: str, inceput: str, sfarsit: str) -> tuple[dict, dict]:
    """Pentru fiecare oră din interval, cea mai RECENTĂ prognoză disponibilă
    (rulările sunt parcurse de la cea mai nouă la cea mai veche, deci prima
    valoare întâlnită e păstrată), plus orizontul la care a fost emisă.

    Orizontul e păstrat pentru că indicatorii calculați pe un amestec de
    orizonturi sunt sistematic mai buni decât realitatea operațională: o
    valoare emisă cu 2 ore înainte nu e comparabilă cu una emisă cu 36."""
    valori: dict[str, float] = {}
    orizonturi: dict[str, float] = {}
    # O rulare declanșată după sfârșitul intervalului prognozează doar zile de
    # după el, iar una mai veche decât `inceput` minus orizontul maxim nu ajunge
    # până în interval — în ambele cazuri nu are nicio oră utilă aici.
    from app import setari as setari_mod
    marja = timedelta(days=int(setari_mod.get(repo, "zile_orizont_maxim")) + 1)
    dupa = (datetime.fromisoformat(inceput) - marja).isoformat()
    for rulare in repo.itereaza_rulari(parc_id, limita=500,
                                        declansata_dupa=dupa, declansata_inainte=sfarsit):
        for ora in rulare.ore:
            if ora.putere_finala_kw is None:
                continue
            if not (inceput <= ora.timestamp_utc <= sfarsit):
                continue
            if ora.timestamp_utc not in valori:
                valori[ora.timestamp_utc] = ora.putere_finala_kw
                if ora.ore_in_avans is not None:
                    orizonturi[ora.timestamp_utc] = ora.ore_in_avans
    return valori, orizonturi


def _serie(dictionar: dict) -> pd.Series:
    if not dictionar:
        return pd.Series(dtype="float64")
    serie = pd.Series(dictionar)
    serie.index = pd.to_datetime(serie.index, utc=True)
    return serie.sort_index()


def genereaza_raport_lunar(repo, parc: Parc, an: int, luna: int, fus_orar: str | None = None,
                            format_timp: str | None = None) -> bytes:
    """`fus_orar` controlează DOAR fusul în care se AFIȘEAZĂ coloana de
    timestamp (implicit — fusul declarat al parcului, `parc.fus_orar`, ales
    de operator la descărcare — R-xxx). Luna calendaristică (`an`, `luna`) și
    filtrarea orelor rămân în UTC, ca înainte — schimbarea fusului de
    afișare nu mută granița lunii și nu poate, deci, să scoată sau să adauge
    ore prin simpla schimbare a fusului ales.

    `format_timp` alege scrierea coloanei de timp, din aceeași listă ca la
    exportul CSV (`app.core.momente.FORMATE_TIMP`) — implicit ISO 8601 cu
    decalaj explicit, ca toate fișierele produse de platformă să arate la fel."""
    from app.core.momente import format_timp_valid, formateaza_moment

    fmt = format_timp_valid(format_timp)
    tz_afisare = ZoneInfo(fus_orar or parc.fus_orar)
    ultima_zi = calendar.monthrange(an, luna)[1]
    inceput = f"{an:04d}-{luna:02d}-01T00:00:00+00:00"
    sfarsit = f"{an:04d}-{luna:02d}-{ultima_zi:02d}T23:59:59+00:00"

    productie = {p.timestamp_utc: p.energie_kwh
                 for p in repo.listeaza_productie(parc.id, inceput, sfarsit)}
    prognoza, orizonturi = _prognoza_pe_ora(repo, parc.id, inceput, sfarsit)

    toate_orele = sorted(set(productie) | set(prognoza))

    wb = Workbook()

    # ------------------------------------------------------------------ date orare
    ws_date = wb.active
    ws_date.title = "Date orare"
    # Coloanele B-E sunt în MW/MWh (cerința operatorului) — calculele interne
    # rămân în kW/kWh, conversia se face O SINGURĂ DATĂ, la scriere. Coloana A
    # arată ora LOCALĂ, în fusul ales la descărcare (implicit fusul parcului),
    # gestionând corect trecerea la ora de vară/iarnă (R-1502); coloana H
    # păstrează timestamp-ul UTC brut — sursa de adevăr, neambiguă indiferent
    # de fusul de afișare ales.
    antete = [f"Timestamp ({fus_orar or parc.fus_orar})", "Prognozat (MW)", "Realizat (MWh)",
              "Eroare absolută (MW)", "Eroare cu semn (MW)", "Ore în avans", "Calitate", "Timestamp (UTC)"]
    for c, text in enumerate(antete, start=1):
        celula = ws_date.cell(row=1, column=c, value=text)
        celula.fill, celula.font = ANTET_FILL, ANTET_FONT

    calitate = {v.timestamp_utc: v.steag for v in repo.listeaza_calitate_productie(parc.id)}

    for i, ts in enumerate(toate_orele, start=2):
        ws_date.cell(row=i, column=1,
                      value=formateaza_moment(ts, fus_orar or parc.fus_orar, fmt))
        ws_date.cell(row=i, column=8, value=ts)
        if ts in prognoza:
            celula = ws_date.cell(row=i, column=2, value=round(kw_in_mw(prognoza[ts]), 3))
            celula.number_format = FORMAT_MW
        if ts in productie:
            celula = ws_date.cell(row=i, column=3, value=round(kw_in_mw(productie[ts]), 3))
            celula.number_format = FORMAT_MW
        # formule REALE (R-1202): editarea coloanelor B/C recalculează totul
        ws_date.cell(row=i, column=4, value=f'=IF(AND(B{i}<>"",C{i}<>""),ABS(B{i}-C{i}),"")').number_format = FORMAT_MW
        ws_date.cell(row=i, column=5, value=f'=IF(AND(B{i}<>"",C{i}<>""),B{i}-C{i},"")').number_format = FORMAT_MW
        if ts in orizonturi:
            ws_date.cell(row=i, column=6, value=round(orizonturi[ts], 1))
        if ts in calitate:
            ws_date.cell(row=i, column=7, value=calitate[ts])

    n_randuri = len(toate_orele)
    ultim_rand = n_randuri + 1
    for col, latime in zip("ABCDEFGH", (17, 16, 16, 18, 18, 13, 20, 24)):
        ws_date.column_dimensions[col].width = latime
    ws_date.freeze_panes = "A2"

    # ------------------------------------------------------------------ sumar
    ws_sumar = wb.create_sheet("Sumar")
    instalata = putere_instalata_kw(parc)
    baza_normalizare = putere_normalizare_kw(parc)

    randuri_sumar: list[tuple[str, object]] = [
        ("Parc", f"{parc.nume} ({parc.cod})"),
        ("Lună", f"{luna:02d}/{an}"),
        ("Putere instalată (MW)", round(kw_in_mw(instalata), 3)),
        ("Bază de normalizare a indicatorilor (MW)", round(kw_in_mw(baza_normalizare), 3)),
        ("Ore cu date", n_randuri),
        ("", ""),
        ("ENERGIE", ""),
        ("Energie totală prognozată (MWh)", f"=SUM('Date orare'!B2:B{ultim_rand})"),
        ("Energie totală realizată (MWh)", f"=SUM('Date orare'!C2:C{ultim_rand})"),
        ("Abatere de energie (%)", '=IFERROR((B8-B9)/B9*100,"")'),
        ("", ""),
        ("ACURATEȚE (normalizată la puterea de referință — convenția operatorilor de sistem)", ""),
        # B4 = baza de normalizare; referințele sunt verificate prin teste automate
        ("Eroare medie absolută — nMAE (% din puterea de referință)",
         f"=IFERROR(AVERAGE('Date orare'!D2:D{ultim_rand})/$B$4*100,\"\")"),
        ("Eroare pătratică medie — nRMSE (% din puterea de referință)",
         f"=IFERROR(SQRT(SUMPRODUCT('Date orare'!E2:E{ultim_rand},'Date orare'!E2:E{ultim_rand})"
         f"/COUNT('Date orare'!E2:E{ultim_rand}))/$B$4*100,\"\")"),
        ("Abatere sistematică — bias (%, cu semn)",
         f"=IFERROR(AVERAGE('Date orare'!E2:E{ultim_rand})/$B$4*100,\"\")"),
        ("", ""),
        ("ACURATEȚE (normalizată la producția medie — convenția traderilor)", ""),
        ("Eroare medie absolută — nMAE (% din producția medie)",
         f"=IFERROR(AVERAGE('Date orare'!D2:D{ultim_rand})/AVERAGE('Date orare'!C2:C{ultim_rand})*100,\"\")"),
        ("", ""),
        ("Fus orar afișat (foaia „Date orare”, coloana A — coloana H păstrează UTC brut)",
         str(tz_afisare)),
    ]
    RANDURI_MW = {3, 4, 8, 9}  # putere instalată, bază normalizare, energie prognozată/realizată
    for i, (eticheta, valoare) in enumerate(randuri_sumar, start=1):
        celula_eticheta = ws_sumar.cell(row=i, column=1, value=eticheta)
        if eticheta and not valoare:
            celula_eticheta.font = SECTIUNE_FONT
        else:
            celula_eticheta.font = Font(bold=True)
        celula_valoare = ws_sumar.cell(row=i, column=2, value=valoare)
        if i in RANDURI_MW:
            celula_valoare.number_format = FORMAT_MW
    ws_sumar.column_dimensions["A"].width = 66
    ws_sumar.column_dimensions["B"].width = 22

    nota = ws_sumar.cell(row=len(randuri_sumar) + 2, column=1, value=(
        "Notă privind interpretarea: indicatorii de mai sus sunt calculați pe TOATE orele "
        "cu date din lună, inclusiv cele de noapte (la fotovoltaic) și cele afectate de "
        "indisponibilități sau limitări de putere. Foaia „Verificare” conține aceiași "
        "indicatori calculați corect — doar pe orele utile și fără orele în care parcul nu a "
        "putut produce din motive care nu țin de prognoză — plus scorul de skill față de o "
        "prognoză de referință."))
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws_sumar.merge_cells(start_row=nota.row, start_column=1, end_row=nota.row + 3, end_column=2)

    # ------------------------------------------------------------------ verificare metodologică
    _adauga_foaie_verificare(wb, repo, parc, _serie(prognoza), _serie(productie),
                              _serie(orizonturi), baza_normalizare)

    # ------------------------------------------------------------------ grafic (R-1204)
    if n_randuri > 1:
        chart = LineChart()
        chart.title = "Prognozat vs. realizat"
        chart.y_axis.title = "MW / MWh"
        chart.x_axis.title = "Oră"
        date_ref = Reference(ws_date, min_col=2, max_col=3, min_row=1, max_row=ultim_rand)
        chart.add_data(date_ref, titles_from_data=True)
        chart.width, chart.height = 26, 12
        ws_date.add_chart(chart, "I2")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _adauga_foaie_verificare(wb, repo, parc: Parc, prognoza: pd.Series, realizat: pd.Series,
                              orizonturi: pd.Series, baza_normalizare: float) -> None:
    """Foaia cu indicatorii calculați conform metodologiei de verificare
    (orele utile, excluderea orelor cu probleme de calitate, scor de skill)."""
    from app.core.clearsky import masca_ore_utile, putere_referinta
    from app.core.data_quality import ore_excluse
    from app.core.verification import verifica

    ws = wb.create_sheet("Verificare")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 48

    def scrie(rand: int, eticheta: str, valoare=None, explicatie: str = "") -> int:
        ws.cell(row=rand, column=1, value=eticheta).font = Font(bold=bool(eticheta and valoare is None))
        if valoare is not None:
            ws.cell(row=rand, column=2, value=valoare)
        if explicatie:
            celula = ws.cell(row=rand, column=3, value=explicatie)
            celula.alignment = Alignment(wrap_text=True, vertical="top")
        return rand + 1

    rand = 1
    rand = scrie(rand, "VERIFICARE METODOLOGICĂ")
    if prognoza.empty or realizat.empty:
        scrie(rand, "Nu există suficiente date comune pentru verificare în această lună.")
        return

    comun = prognoza.index.intersection(realizat.index)
    referinta_cs = putere_referinta(parc, comun) if len(comun) else None
    utile = masca_ore_utile(parc, referinta_cs, baza_normalizare) if referinta_cs is not None else None
    excluse = ore_excluse(repo, parc.id)

    raport = verifica(
        prognozat_kw=prognoza, realizat_kw=realizat, putere_normalizare_kw=baza_normalizare,
        referinta_cer_senin_kw=referinta_cs, masca_ore_utile=utile,
        ore_excluse_qc=excluse if len(excluse) else None,
        ore_in_avans=orizonturi if not orizonturi.empty else None,
    )
    ind = raport.indicatori

    rand = scrie(rand, "Ore incluse în calcul", ind.numar_ore,
                  "Doar orele utile (la fotovoltaic: orele de zi) și fără orele marcate de "
                  "controlul de calitate.")
    rand = scrie(rand, "Ore excluse — noapte / fără producție posibilă", raport.ore_excluse_noapte,
                  "Includerea lor ar face indicatorul artificial de bun: zerourile de noapte se "
                  "„prognozează” perfect.")
    rand = scrie(rand, "Ore excluse — probleme de calitate", raport.ore_excluse_qc,
                  "Indisponibilități, limitări de putere, valori blocate — măsoară starea "
                  "parcului, nu calitatea prognozei.")
    rand += 1
    rand = scrie(rand, "INDICATORI (bază de normalizare: puterea de referință a parcului)")
    rand = scrie(rand, "nMAE (%)", _rotunjit(ind.nmae_capacitate_pct))
    rand = scrie(rand, "nRMSE (%)", _rotunjit(ind.nrmse_capacitate_pct))
    rand = scrie(rand, "Bias (%, cu semn)", _rotunjit(ind.nbias_capacitate_pct),
                  "Pozitiv = supraestimare sistematică. Un bias persistent e cel mai bun "
                  "indiciu al unei limitări de putere neanunțate sau al unei configurații greșite.")
    rand = scrie(rand, "nMAE raportat la producția medie (%)", _rotunjit(ind.nmae_medie_pct),
                  "Convenția folosită de traderi; mereu mai mare decât cel normalizat la capacitate.")
    rand += 1
    rand = scrie(rand, "SCOR DE SKILL FAȚĂ DE PROGNOZA DE REFERINȚĂ")
    rand = scrie(rand, "Skill (0 = cât referința, 1 = perfect)", _rotunjit(raport.skill_vs_referinta, 3),
                  "Referința e combinația optimă persistență/climatologie (la fotovoltaic, în "
                  "spațiul indicelui de senin). Un scor ≤ 0 înseamnă că modelul nu aduce nimic "
                  "peste o regulă simplă.")
    rand = scrie(rand, "nMAE al prognozei de referință (%)",
                  _rotunjit(raport.indicatori_referinta.nmae_capacitate_pct))

    if raport.pe_orizont:
        rand += 1
        rand = scrie(rand, "INDICATORI PE ORIZONT DE PROGNOZĂ")
        for eticheta, indicatori in raport.pe_orizont.items():
            rand = scrie(rand, f"  {eticheta} — nMAE (%)", _rotunjit(indicatori.nmae_capacitate_pct),
                          f"{indicatori.numar_ore} ore")

    if raport.crps_kw is not None:
        rand += 1
        rand = scrie(rand, "VERIFICARE PROBABILISTICĂ")
        rand = scrie(rand, "CRPS (MW)", _rotunjit(kw_in_mw(raport.crps_kw), 3),
                      "Scorul propriu pentru prognoza probabilistică; mai mic = mai bun.")


def _rotunjit(valoare, zecimale: int = 2):
    return None if valoare is None else round(float(valoare), zecimale)


def genereaza_raport_verificare(repo, parc: Parc, zile: int = 90) -> dict:
    """Raportul de verificare afișat în interfață (nu în Excel): indicatori,
    skill, defalcare pe orizont, pentru ultimele `zile` zile."""
    from app.core.clearsky import masca_ore_utile, putere_referinta
    from app.core.data_quality import ore_excluse
    from app.core.verification import verifica

    sfarsit = datetime.now(timezone.utc)
    inceput = sfarsit - pd.Timedelta(days=zile)
    inceput_text, sfarsit_text = iso_utc(inceput), iso_utc(sfarsit)
    productie = {p.timestamp_utc: p.energie_kwh
                 for p in repo.listeaza_productie(parc.id, inceput_text, sfarsit_text)}
    prognoza, orizonturi = _prognoza_pe_ora(repo, parc.id, inceput_text, sfarsit_text)

    serie_prognoza, serie_realizat = _serie(prognoza), _serie(productie)
    if serie_prognoza.empty or serie_realizat.empty:
        return {"disponibil": False, "zile": zile}

    comun = serie_prognoza.index.intersection(serie_realizat.index)
    if len(comun) == 0:
        return {"disponibil": False, "zile": zile}

    baza = putere_normalizare_kw(parc)
    referinta_cs = putere_referinta(parc, comun)
    raport = verifica(
        prognozat_kw=serie_prognoza, realizat_kw=serie_realizat, putere_normalizare_kw=baza,
        referinta_cer_senin_kw=referinta_cs,
        masca_ore_utile=masca_ore_utile(parc, referinta_cs, baza),
        ore_excluse_qc=ore_excluse(repo, parc.id),
        ore_in_avans=_serie(orizonturi) if orizonturi else None,
    )
    return {"disponibil": True, "zile": zile, "raport": raport, "baza_normalizare_kw": baza}
