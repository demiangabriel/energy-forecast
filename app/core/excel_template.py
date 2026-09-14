"""
Șabloane Excel descărcabile pentru introducerea în masă a parametrilor unui
parc nou — R-407, Anexele B și C. Fiecare coloană are un comentariu Excel cu
explicația câmpului (importanță + ce influențează), ca operatorul să poată
pregăti datele offline, fără să deschidă aplicația.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

ANTET_FILL = PatternFill(start_color="1F6F50", end_color="1F6F50", fill_type="solid")
ANTET_FONT = Font(color="FFFFFF", bold=True)

COLOANE_PV = [
    ("cod_parc", "important", "Identificator unic, scurt. Nu se schimbă după creare."),
    ("nume", "important", "Identificare administrativă."),
    ("client", "important", "Numele clientului/organizației proprietare."),
    ("latitudine", "important", "Coordonata exactă a parcului (grade zecimale)."),
    ("longitudine", "important", "Coordonata exactă a parcului (grade zecimale)."),
    ("altitudine_m", "opțional", "Impact sub 0,1% pe rezultat."),
    ("fus_orar", "important", "Orice fus orar IANA valid (ex. Europe/Bucharest, Europe/Chisinau, "
     "Europe/Sofia, Europe/Budapest) — vezi lista completă în aplicație, la adăugarea parcului."),
    ("data_punerii_in_functiune", "opțional", "Folosită pentru calculul degradării (AAAA-LL-ZZ)."),
    ("putere_dc_kwp", "important", "Suma tuturor panourilor (kWp)."),
    ("putere_ac_kw", "important", "Limita de vârf a parcului / puterea invertorului (kW)."),
    ("randament_invertor", "opțional", "Implicit 0,96-0,98."),
    ("limita_injectie_kw", "opțional", "Completați doar dacă avizul de racordare o impune."),
    ("tehnologie_panou", "opțional", "monocrystalline / polycrystalline / PERC_TOPCon_HJT / thin_film / bifacial."),
    ("inclinare_grade", "opțional*", "Poate fi dedusă ulterior prin calibrare (secțiunea 7) dacă lipsește."),
    ("azimut_grade", "opțional*", "180 = sud. Poate fi dedusă prin calibrare dacă lipsește."),
    ("tip_asezare", "important", "fixed / east_west / tracker — impact 20-40% pe producție."),
    ("coeficient_temperatura_pct_per_c", "opțional", "Implicit, în funcție de tehnologia panoului."),
    ("model_termic_montaj", "opțional", "open_rack_glass_glass / open_rack_glass_polymer / insulated_back_glass_polymer."),
    ("degradare_anuala_pct", "opțional", "Implicit 0,4-0,8%/an, în funcție de vechime."),
    ("pierdere_murdarire_pct", "opțional", "Compun multiplicativ cu celelalte pierderi de sistem."),
    ("pierdere_neuniformitate_pct", "opțional", "Compun multiplicativ."),
    ("pierdere_cablaj_conexiuni_pct", "opțional", "Compun multiplicativ."),
    ("pierdere_lid_pct", "opțional", "Compun multiplicativ."),
    ("pierdere_placuta_pct", "opțional", "Compun multiplicativ."),
    ("pierdere_indisponibilitate_pct", "opțional", "Compun multiplicativ."),
    ("umbrire_orizont_grade", "opțional", "Gradul unghiular sub care soarele e blocat de orizont."),
    ("umbrire_orizont_directie", "opțional", "est / vest / ambele."),
]

COLOANE_EOLIAN = [
    ("cod_parc", "important", "Identificator unic, scurt. Nu se schimbă după creare."),
    ("nume", "important", "Identificare administrativă."),
    ("client", "important", "Numele clientului/organizației proprietare."),
    ("latitudine", "important", "Coordonata exactă a parcului (grade zecimale)."),
    ("longitudine", "important", "Coordonata exactă a parcului (grade zecimale)."),
    ("altitudine_m", "opțional", "Impact sub 0,1% pe rezultat."),
    ("fus_orar", "important", "Orice fus orar IANA valid (ex. Europe/Bucharest, Europe/Chisinau, "
     "Europe/Sofia, Europe/Budapest) — vezi lista completă în aplicație, la adăugarea parcului."),
    ("numar_turbine", "important", ""),
    ("putere_nominala_kw", "important", "Putere nominală PER turbină (kW)."),
    ("inaltime_butuc_m", "important", "Vântul crește mult cu înălțimea."),
    ("curba_putere_fisier", "opțional", "Lăsați gol dacă completați vitezele simple mai jos. "
                                         "Dacă aveți fișa tehnică, introduceți curba separat, în aplicație."),
    ("viteza_pornire_ms", "opțional", "Folosită dacă nu există curba exactă a producătorului."),
    ("viteza_nominala_ms", "opțional", "Folosită dacă nu există curba exactă a producătorului."),
    ("viteza_oprire_ms", "opțional", "Folosită dacă nu există curba exactă a producătorului."),
    ("pierdere_siaj_pct", "opțional", "Implicit ~8%; mai mult la parcuri dese."),
    ("pierdere_sistem_pct", "opțional", ""),
]


def _construieste(coloane: list[tuple[str, str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Parc nou"
    for col_idx, (nume, importanta, explicatie) in enumerate(coloane, start=1):
        celula = ws.cell(row=1, column=col_idx, value=nume)
        celula.fill = ANTET_FILL
        celula.font = ANTET_FONT
        text_comentariu = f"[{importanta}] {explicatie}" if explicatie else f"[{importanta}]"
        celula.comment = Comment(text_comentariu, "Platformă prognoză")
        ws.column_dimensions[celula.column_letter].width = max(18, len(nume) + 2)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sablon_pv() -> bytes:
    return _construieste(COLOANE_PV)


def sablon_eolian() -> bytes:
    return _construieste(COLOANE_EOLIAN)
